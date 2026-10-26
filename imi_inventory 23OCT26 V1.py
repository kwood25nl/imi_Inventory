"""IMI Inventory Tracker - Find Item + Create Item.

Data lives in a hidden, password-protected Excel file next to this program.
Build to an exe:  pyinstaller --onefile --noconsole imi_inventory.py
"""
import contextlib
import datetime as dt
import io
import os
import re
import sys
import time
import tkinter as tk
import tkinter.font as tkfont
from tkinter import messagebox, ttk

import msoffcrypto
import openpyxl

# ----------------------------------------------------------------- config
APP_DIR = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
DATA_FILE = os.path.join(APP_DIR, "imi_inventory_data.xlsx")
LOCK_FILE = DATA_FILE + ".lock"
PASSWORD = os.environ.get("IMI_DATA_PASSWORD", "change-me")  # TODO: set your own
BUSY_MSG = ("Cannot perform action. Base excel file is open by one of your colleagues. "
            "Please ask them to close this file to complete action.")
CATEGORIES = ["00 External", "01 PCB", "02 PCBA", "03 Cable/Harness", "04 Mechanical Part",
              "05 Assembly", "06 Chemical Formulation", "07 Firmware/Software",
              "08 Purchased Component (no MPN)", "09 Other", "10 Documents"]
DEFAULT_UOMS = ["Unit", "Vials", "Grams (g)", "Milligrams (mg)", "Kilogram (kg)",
                "Liters (L)", "Milliliters (mL)", "Molar mass (KU)"]
DISCRETE_UOMS = {"unit", "units", "vial", "vials"}
SUBCATS = {"REAG": "Commercial reagents", "ACID": "Acids", "BASE": "Bases", "SALT": "Salts",
           "BUBL": "Synthesized building blocks", "FINL": "Final, synthesized, pure products",
           "SOLV": "Solvents", "MISC": "Miscellaneous"}
REV_TRACKED = {"01", "02", "03", "04", "05", "06", "10"}  # revision always starts at A
REV_OPTIONAL = {"08", "09"}                               # "as needed"
DASH_CATS = {"02", "03", "04", "05"}
REV_LETTERS = "ABCDEFGHJKLMNPRSTUVWXYZ"  # no I, O, Q
REV_SEQ = list(REV_LETTERS) + [a + b for a in REV_LETTERS for b in REV_LETTERS]  # A..Z, AA, AB, ...
# NN[-AAAA]-NNNNN[A or AA, never I/O/Q][-NN]
NUM_RE = re.compile(r"^(?P<cat>\d{2})(?:-(?P<sub>[A-Z]{4}))?-(?P<seq>\d{5})(?P<rev>[A-HJ-NPR-Z]{1,2})?(?:-(?P<dash>\d{2}))?$")
SCHEMA = {
    "Items": ["item_number", "name", "category", "quantity", "uom", "location", "manufacturer",
              "notes", "tracking", "low_stock_alert", "min_quantity",
              "created_at", "created_by", "edited_at", "edited_by", "subcategory", "revision", "dash_number", "status"],
    "Counters": ["category", "last_sequence"],
    "Users": ["initials", "first_name", "last_name", "email", "registered_at"],
    "UoM": ["name"],
    "Locations": ["name"],
    "Lots": ["item_number", "lot_code", "quantity", "uom", "location", "created_at"],
    "Activity": ["timestamp", "initials", "message"],
}
FMT = "%Y-%m-%d %H:%M:%S"


class FileBusy(Exception):
    pass


# ---------------------------------------------------------------- storage
def _norm(s):
    return re.sub(r"\s+", "", str(s)).casefold()


def build_number(code, seq, sub="", rev="", dash=""):
    return f"{code}" + (f"-{sub}" if sub else "") + f"-{int(seq):05d}" + (rev or "") + (f"-{dash}" if dash else "")


def parse_number(text):
    """Validate a part number against the IMI Part Numbering System. Returns its fields."""
    t = re.sub(r"\s+", "", str(text)).upper()
    m = NUM_RE.match(t)
    if not m:
        raise ValueError(f"{text} doesn't match the IMI part number format NN-AAAA-NNNNNA-NN "
                         "(e.g. 01-00042A, 02-00042A-01, 06-REAG-00001A).")
    d = m.groupdict()
    code = d["cat"]
    if code not in {c[:2] for c in CATEGORIES} or code == "00":
        raise ValueError(f"{code} is not a valid category code.")
    if code == "06":
        if d["sub"] not in SUBCATS:
            raise ValueError("Chemical formulations (06) need a subcategory: " + ", ".join(SUBCATS) + ".")
    elif d["sub"]:
        raise ValueError(f"Category {code} doesn't use a subcategory.")
    if d["dash"] and code not in DASH_CATS:
        raise ValueError(f"Dash numbers are only used for categories {', '.join(sorted(DASH_CATS))}.")
    if code == "07" and d["rev"]:
        raise ValueError("Firmware/Software (07) doesn't track revisions here (GitLab version control).")
    if code in REV_TRACKED and not d["rev"]:
        raise ValueError(f"Category {code} tracks revisions: add a revision letter (first release is A).")
    if int(d["seq"]) == 0:
        raise ValueError("The sequence starts at 00001.")
    d["cat_code"] = code
    return d


def _hide(path):
    if os.name == "nt":
        import ctypes
        ctypes.windll.kernel32.SetFileAttributesW(path, 0x02)  # FILE_ATTRIBUTE_HIDDEN


def _excel_owner_file_exists():
    d, n = os.path.split(DATA_FILE)
    return any(os.path.exists(os.path.join(d, p)) for p in ("~$" + n, "~$" + n[2:]))


def _acquire_lock():
    for _ in range(25):
        try:
            os.close(os.open(LOCK_FILE, os.O_CREAT | os.O_EXCL | os.O_WRONLY))
            return
        except FileExistsError:
            with contextlib.suppress(OSError):
                if time.time() - os.path.getmtime(LOCK_FILE) > 30:  # stale lock
                    os.remove(LOCK_FILE)
            time.sleep(0.2)
    raise FileBusy


def _encrypt(wb):
    plain = io.BytesIO()
    wb.save(plain)
    plain.seek(0)
    out = io.BytesIO()
    msoffcrypto.OfficeFile(plain).encrypt(PASSWORD, out)
    return out.getvalue()


def _create_file():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in SCHEMA.items():
        wb.create_sheet(name).append(cols)
    for u in DEFAULT_UOMS:
        wb["UoM"].append([u])
    wb["Locations"].append(["Virtual Intake"])
    with open(DATA_FILE, "wb") as f:
        f.write(_encrypt(wb))
    _hide(DATA_FILE)


@contextlib.contextmanager
def transaction(write=False):
    """Open the workbook. Writers are blocked if a colleague has it open in Excel."""
    if write:
        if _excel_owner_file_exists():
            raise FileBusy
        _acquire_lock()
    try:
        if not os.path.exists(DATA_FILE):
            _create_file()
        try:
            fh = open(DATA_FILE, "r+b" if write else "rb")  # Excel holds a write-deny lock
        except PermissionError:
            raise FileBusy
        with fh:
            plain = io.BytesIO()
            office = msoffcrypto.OfficeFile(fh)
            office.load_key(password=PASSWORD)
            office.decrypt(plain)
            wb = openpyxl.load_workbook(plain)
            _ensure(wb)
            yield wb
            if write:
                data = _encrypt(wb)
                fh.seek(0)
                fh.truncate()
                fh.write(data)
    finally:
        if write:
            with contextlib.suppress(OSError):
                os.remove(LOCK_FILE)


def _ensure(wb):
    """Add any sheets/columns introduced by newer versions to an existing file."""
    for name, cols in SCHEMA.items():
        if name not in wb.sheetnames:
            wb.create_sheet(name).append(cols)
            continue
        ws = wb[name]
        have = [c.value for c in ws[1]]
        for c in cols:
            if c not in have:
                have.append(c)
                ws.cell(row=1, column=len(have), value=c)


def _append(ws, rec):
    ws.append([rec.get(c.value) for c in ws[1]])


def status_of(item):
    return item.get("status") or "Active"


def fmt_ts(text):
    try:
        return dt.datetime.strptime(text, FMT).strftime("%b %d, %Y %I:%M %p")
    except (TypeError, ValueError):
        return str(text or "")


def _row_index(ws, number):
    head = [c.value for c in ws[1]]
    col = head.index("item_number") + 1
    for r in range(2, ws.max_row + 1):
        if _norm(ws.cell(r, col).value or "") == _norm(number):
            return r, head
    raise ValueError(f"Item {number} was not found.")


def _set_cells(ws, r, head, **vals):
    for k, val in vals.items():
        ws.cell(r, head.index(k) + 1).value = val


def _rows(ws):
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def _now():
    return dt.datetime.now().strftime(FMT)


def _log(wb, initials, msg):
    wb["Activity"].append([_now(), initials, msg])


def _pick_sequence(wb, code, items):
    """Next unused sequence for a category: counter + 1, skipping any number already taken."""
    cats = {"01", "02"} if code in ("01", "02") else {code}  # PCB / PCBA share numbers only when linked
    used = {int(m["seq"]) for i in items if (m := NUM_RE.match(str(i["item_number"]))) and m["cat"] in cats}
    row = next((r for r in wb["Counters"].iter_rows(min_row=2) if r[0].value == code), None)
    seq = (int(row[1].value) if row else 0) + 1
    while seq in used:
        seq += 1
    return seq


def _seqs(items, code):
    """sequence -> first item number using it, for one category."""
    out = {}
    for i in items:
        m = NUM_RE.match(str(i["item_number"]))
        if m and m["cat"] == code:
            out.setdefault(int(m["seq"]), i["item_number"])
    return out


def _set_counter(wb, code, seq):
    row = next((r for r in wb["Counters"].iter_rows(min_row=2) if r[0].value == code), None)
    if row:
        row[1].value = seq
    else:
        wb["Counters"].append([code, seq])


class DB:
    def lookups(self):
        with transaction() as wb:
            return ([r["name"] for r in _rows(wb["UoM"])], [r["name"] for r in _rows(wb["Locations"])])

    def items(self):
        with transaction() as wb:
            return _rows(wb["Items"])

    def activity(self, n=50):
        with transaction() as wb:
            return _rows(wb["Activity"])[-n:][::-1]

    def peek_sequence(self, code):
        with transaction() as wb:
            return _pick_sequence(wb, code, _rows(wb["Items"]))

    def find_data(self):
        """All items plus {item_number: [lot codes]} for the Find window and item cards."""
        with transaction() as wb:
            lots = {}
            for l in _rows(wb["Lots"]):
                lots.setdefault(l["item_number"], []).append(str(l["lot_code"]))
            return _rows(wb["Items"]), lots

    def update_item(self, number, changes, user):
        with transaction(write=True) as wb:
            ws = wb["Items"]
            r, head = _row_index(ws, number)
            _set_cells(ws, r, head, **changes, edited_at=_now(), edited_by=user)
            _log(wb, user, f"Edited item: {number}")
        return True

    def set_status(self, number, status, user):
        with transaction(write=True) as wb:
            ws = wb["Items"]
            r, head = _row_index(ws, number)
            _set_cells(ws, r, head, status=status, edited_at=_now(), edited_by=user)
            _log(wb, user, f"{'Archived' if status == 'Archived' else 'Reactivated'} item: {number}")
        return True

    def roll_revision(self, number, new_rev, mode, user):
        """mode 'convert': rename the item (all stock moves to the new rev).
        mode 'new': create the new revision as a separate item; the old rev keeps its stock."""
        with transaction(write=True) as wb:
            ws = wb["Items"]
            r, head = _row_index(ws, number)
            d = NUM_RE.match(str(number))
            if not d or d["cat"] not in REV_TRACKED | REV_OPTIONAL:
                raise ValueError("This item doesn't use revisions.")
            new = build_number(d["cat"], d["seq"], d["sub"] or "", new_rev, d["dash"] or "")
            if any(_norm(i["item_number"]) == _norm(new) for i in _rows(ws)):
                raise ValueError(f"{new} already exists.")
            now = _now()
            if mode == "convert":
                _set_cells(ws, r, head, item_number=new, revision=new_rev, edited_at=now, edited_by=user)
                lw = wb["Lots"]
                for lr in range(2, lw.max_row + 1):
                    if lw.cell(lr, 1).value == number:
                        lw.cell(lr, 1).value = new
                _log(wb, user, f"Rolled revision: {number} to {new}, all stock converted")
            else:
                row = {h: ws.cell(r, i + 1).value for i, h in enumerate(head)}
                row.update(item_number=new, revision=new_rev, quantity=None, uom=None, location=None,
                           created_at=now, created_by=user, edited_at=now, edited_by=user, status="Active")
                _append(ws, row)
                _log(wb, user, f"Created new revision: {new} (from {number})")
            return new

    def find_user(self, initials):
        with transaction() as wb:
            return next((u for u in _rows(wb["Users"]) if u["initials"].upper() == initials.upper()), None)

    def add_user(self, initials, first, last, email):
        """Returns the initials actually assigned (a number is appended if already taken)."""
        with transaction(write=True) as wb:
            taken = {str(u["initials"]).upper() for u in _rows(wb["Users"])}
            final, n = initials.upper(), 2
            while final in taken:
                final = f"{initials.upper()}{n}"
                n += 1
            wb["Users"].append([final, first, last, email, _now()])
            _log(wb, final, f"New user registered: {first} {last}")
            return final

    def save_item(self, rec, lots, user, auto_number):
        """Returns the final item number. Raises ValueError for invalid or duplicate numbers."""
        with transaction(write=True) as wb:
            items = _rows(wb["Items"])
            code = rec["category"][:2]
            if code == "00":  # external numbers (MPN / vendor): free text
                num = re.sub(r"\s+", " ", str(rec["item_number"])).strip()
                if not num:
                    raise ValueError("Item number is required.")
                rec.update(item_number=num, subcategory=None, revision=None, dash_number=None)
            else:
                linked = bool(rec.get("link_seq"))
                if auto_number:  # picked under the file lock so two people can't collide
                    seq = rec["link_seq"] if linked else _pick_sequence(wb, code, items)
                    num = build_number(code, seq, rec.get("subcategory"), rec.get("revision"), rec.get("dash_number"))
                else:
                    num = re.sub(r"\s+", "", str(rec["item_number"])).upper()
                d = parse_number(num)
                if d["cat_code"] != code:
                    raise ValueError(f"The item number starts with {d['cat_code']} but the category is {rec['category']}.")
                rec.update(item_number=num, subcategory=d["sub"], revision=d["rev"], dash_number=d["dash"])
                if code in ("01", "02"):
                    other = "02" if code == "01" else "01"
                    oseqs, n_ = _seqs(items, other), int(d["seq"])
                    if linked and n_ != int(rec["link_seq"]):
                        raise ValueError(f"The item number's sequence must match the linked part ({int(rec['link_seq']):05d}).")
                    if linked and n_ not in oseqs:
                        raise ValueError(f"No {other} part with sequence {n_:05d} exists to link to.")
                    if not linked and n_ in oseqs and n_ not in _seqs(items, code):
                        raise ValueError(f"Sequence {n_:05d} is already used by {oseqs[n_]}. A 01 PCB and a 02 PCBA "
                                         "can only share a number if they are linked: select it under Linked.")
                if auto_number and not linked:
                    _set_counter(wb, code, seq)
            num = rec["item_number"]
            if any(_norm(i["item_number"]) == _norm(num) for i in items):
                raise ValueError(f"Item number {num} already exists. Every item must have a unique number.")
            if rec["uom"] and not any(_norm(u["name"]) == _norm(rec["uom"]) for u in _rows(wb["UoM"])):
                wb["UoM"].append([rec["uom"]])
            now = _now()
            rec.update(created_at=now, created_by=user, edited_at=now, edited_by=user, status="Active")
            _append(wb["Items"], rec)
            for lot in lots:
                wb["Lots"].append([num, lot[0], lot[1], rec["uom"], rec["location"], now])
            _log(wb, user, f"Created item: {num} {rec['name']}")
            return num


# ---------------------------------------------------------- lot generator
def make_lots(qty, uom, count_text, today=None):
    """YYJJJ lot codes; '-n' suffix when more than one. Returns [(code, qty)]."""
    today = today or dt.date.today()
    base = f"{today:%y}{today.timetuple().tm_yday:03d}"
    discrete = _norm(uom) in DISCRETE_UOMS
    txt = count_text.strip()
    if txt.upper() in ("", "ALL"):
        if discrete and qty != int(qty):
            raise ValueError("Discrete quantities must be whole numbers to generate lot codes.")
        n = int(qty) if discrete else 1
    else:
        if not txt.isdigit() or int(txt) < 1:
            raise ValueError("Lot code count must be a whole number of at least 1, or ALL.")
        n = int(txt)
    if n > qty:
        raise ValueError("The number of lot codes cannot exceed the quantity.")
    codes = [base] if n == 1 else [f"{base}-{i}" for i in range(1, n + 1)]
    if discrete:
        q, r = divmod(int(qty), n)
        shares = [q + (1 if i < r else 0) for i in range(n)]
    else:
        shares = [round(qty / n, 6)] * n
    return list(zip(codes, shares))


# --------------------------------------------------------------------- UI
BG, HEADER, PANEL, PANEL2 = "#14151C", "#1E2029", "#1B1D27", "#262938"
FIELD, BORDER, TXT, MUTED = "#14151C", "#3A3E52", "#F2F3F8", "#A3A9B8"
DIS_BG, DIS_TX, DIS_BD = "#1D1F2A", "#9298A9", "#2E3140"
YELLOW, ACCENT, FOCUS, ERR = "#FFC800", "#7B68EE", "#B8ADFF", "#FF7A7A"
FONT, BOLD = ("Segoe UI", 11), ("Segoe UI", 11, "bold")
INI_COLORS = ["#A99CFF", "#FF8FC0", "#6FD6FA", "#FFD54A", "#8EE07A", "#FFA95E"]
ACTIONS = [("Find Item", "#49CCF9"), ("Create Item", "#6BC950"),
           ("Adjust/Move Inventory", "#FD71AF"), ("Build Mixture/Assembly", "#FF8A3D")]


def blend(fg, bg, a):
    f, b = [tuple(int(c[i:i + 2], 16) for i in (1, 3, 5)) for c in (fg, bg)]
    return "#%02x%02x%02x" % tuple(round(x * a + y * (1 - a)) for x, y in zip(f, b))


def mk_entry(parent, var=None, **kw):
    return tk.Entry(parent, textvariable=var, bg=FIELD, fg=TXT, insertbackground=TXT, relief="flat", font=FONT,
                    highlightthickness=2, highlightbackground=BORDER, highlightcolor=FOCUS,
                    disabledbackground=DIS_BG, disabledforeground=DIS_TX, **kw)


def mk_button(parent, text, cmd, primary=False):
    return tk.Button(parent, text=text, command=cmd, relief="flat", bd=0, font=BOLD, padx=22, pady=8,
                     bg="#6BC950" if primary else PANEL2, fg="#10120F" if primary else TXT,
                     activebackground="#31355A", activeforeground=TXT, cursor="hand2")


class Toggle(tk.Canvas):
    """Square on/off switch bound to a BooleanVar."""
    def __init__(self, master, var, command=None, bg=PANEL):
        super().__init__(master, width=44, height=22, bg=bg, highlightthickness=0, cursor="hand2")
        self.var, self.cmd, self.enabled = var, command, True
        self.bind("<Button-1>", self._click)
        var.trace_add("write", lambda *a: self._draw())
        self._draw()

    def _draw(self):
        self.delete("all")
        on = self.var.get()
        track = (ACCENT if on else BORDER) if self.enabled else DIS_BD
        self.create_rectangle(0, 0, 44, 22, fill=track, outline="")
        x = 24 if on else 2
        self.create_rectangle(x, 2, x + 18, 20, fill=TXT if self.enabled else "#555B70", outline="")

    def set_enabled(self, e):
        self.enabled = e
        self.configure(cursor="hand2" if e else "arrow")
        self._draw()

    def _click(self, _):
        if self.enabled:
            self.var.set(not self.var.get())
            if self.cmd:
                self.cmd()


class RegisterDialog(tk.Toplevel):
    def __init__(self, master, initials=None):
        super().__init__(master, bg=PANEL, padx=28, pady=22)
        self.title("New user")
        heading = f"First time using “{initials}”" if initials else "Switch user"
        sub = ("Please tell us who you are. You'll only need your initials next time." if initials else
               "Enter the new user's details. If the initials are taken, a number is added.")
        self.transient(master)
        self.resizable(False, False)
        self.result = None
        self.fixed = initials
        tk.Label(self, text=heading, font=("Segoe UI", 14, "bold"), bg=PANEL, fg=TXT).grid(columnspan=2, sticky="w")
        tk.Label(self, text=sub, bg=PANEL, fg=MUTED, font=FONT).grid(columnspan=2, sticky="w", pady=(0, 14))
        self.vars = {}
        fields = [("first", "First name"), ("last", "Last name"), ("email", "Email")]
        if not initials:
            fields.insert(0, ("initials", "Initials"))
        for r, (key, label) in enumerate(fields, 2):
            tk.Label(self, text=label, bg=PANEL, fg=TXT, font=BOLD).grid(row=r, column=0, sticky="w", pady=5)
            self.vars[key] = tk.StringVar()
            mk_entry(self, self.vars[key], width=30).grid(row=r, column=1, padx=(14, 0), ipady=5)
        self.err = tk.Label(self, text="", bg=PANEL, fg=ERR, font=FONT)
        self.err.grid(columnspan=2, sticky="w", pady=(8, 0))
        bar = tk.Frame(self, bg=PANEL)
        bar.grid(columnspan=2, sticky="e", pady=(10, 0))
        mk_button(bar, "Cancel", self.destroy).pack(side="left", padx=8)
        mk_button(bar, "Register", self._ok, primary=True).pack(side="left")
        self.bind("<Return>", lambda e: self._ok())
        self.grab_set()
        self.wait_visibility()
        self.focus_force()

    def _ok(self):
        v = {k: x.get().strip() for k, x in self.vars.items()}
        if self.fixed:
            v["initials"] = self.fixed
        if not all(v.values()):
            self.err.config(text="All fields are required.")
        elif not re.fullmatch(r"[A-Za-z0-9]{1,4}", v["initials"]):
            self.err.config(text="Initials must be 1 to 4 letters or numbers.")
        elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", v["email"]):
            self.err.config(text="Enter a valid email address.")
        else:
            self.result = v
            self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("IMI Inventory Tracker")
        self.geometry("1280x800")
        self.minsize(1100, 720)
        self.configure(bg=BG)
        self.db = DB()
        self.user = None
        self.content_widget = None
        self._style()
        self._header()
        self._actions()
        self.content = tk.Frame(self, bg=PANEL, highlightthickness=1, highlightbackground=DIS_BD)
        self.content.pack(fill="both", expand=True, padx=32, pady=20)
        self._activity()
        self._tick()
        self.show_find()
        self.refresh_activity()

    # -- shared plumbing
    def guarded(self, fn, *a, **kw):
        """Run a DB call; show the colleague-has-it-open popup on lock errors."""
        try:
            return fn(*a, **kw)
        except FileBusy:
            messagebox.showerror("Error", BUSY_MSG)
        except Exception as e:  # wrong password, corrupt file, etc.
            messagebox.showerror("Error", f"Could not access the data file.\n\n{e}")
        return None

    def _style(self):
        self.option_add("*TCombobox*Listbox.background", FIELD)
        self.option_add("*TCombobox*Listbox.foreground", TXT)
        self.option_add("*TCombobox*Listbox.selectBackground", "#31355A")
        self.option_add("*TCombobox*Listbox.font", FONT)
        s = ttk.Style()
        s.theme_use("clam")
        s.configure("TCombobox", fieldbackground=FIELD, background=PANEL2, foreground=TXT, arrowcolor=TXT,
                    bordercolor=BORDER, lightcolor=BORDER, darkcolor=BORDER, padding=6)
        s.map("TCombobox", fieldbackground=[("readonly", FIELD), ("disabled", DIS_BG)],
              foreground=[("disabled", DIS_TX)], selectbackground=[("readonly", FIELD)],
              selectforeground=[("readonly", TXT)], bordercolor=[("focus", FOCUS)])
        s.configure("Treeview", background=PANEL, fieldbackground=PANEL, foreground=TXT, rowheight=32,
                    borderwidth=0, font=FONT)
        s.configure("Treeview.Heading", background=PANEL2, foreground=TXT, relief="flat", font=BOLD, padding=6)
        s.map("Treeview", background=[("selected", "#31355A")])
        s.configure("Vertical.TScrollbar", background="#6A7088", troughcolor="#23263A", bordercolor="#23263A",
                    arrowcolor=TXT)

    def _header(self):
        box = tk.Frame(self, bg=HEADER)
        box.pack(fill="x")
        row = tk.Frame(box, bg=HEADER)
        row.pack(fill="x", padx=32, pady=20)
        row.columnconfigure((0, 2), weight=1, uniform="h")
        left = tk.Frame(row, bg=HEADER)
        left.grid(row=0, column=0, sticky="w")
        tk.Label(left, text="Enter Initials", bg=HEADER, fg=TXT, font=("Segoe UI", 12, "bold")).pack(side="left")
        self.ini_var = tk.StringVar()
        self.ini = tk.Entry(left, textvariable=self.ini_var, width=6, bg=BG, fg=TXT, insertbackground=TXT,
                            relief="flat", font=("Segoe UI", 16, "bold"), highlightthickness=2,
                            highlightbackground=YELLOW, highlightcolor=FOCUS)
        self.ini.pack(side="left", padx=12, ipady=5)
        self.ini_var.trace_add("write", lambda *a: self.ini_var.set(self.ini_var.get().upper()[:6]))
        self.ini.bind("<Return>", lambda e: self.sign_in())
        self.ini.bind("<FocusOut>", lambda e: self.sign_in())
        self.chip = tk.Label(left, text="🔒 Required to create or edit data", bg="#3A3115", fg="#FFD54A",
                             font=("Segoe UI", 9, "bold"), padx=10, pady=4)
        who = tk.Frame(left, bg=HEADER)
        who.pack(side="left")
        self.chip.pack(in_=who, anchor="w")
        self.chip.lift()
        self.not_link = tk.Label(who, text="", bg=HEADER, fg=FOCUS, cursor="hand2",
                                 font=("Segoe UI", 9, "underline"))
        self.not_link.pack(anchor="w", pady=(3, 0))
        self.not_link.bind("<Button-1>", lambda e: self.switch_user())
        mid = tk.Frame(row, bg=HEADER)
        mid.grid(row=0, column=1, padx=24)
        self.clock = tk.Label(mid, bg=HEADER, fg=TXT, font=("Segoe UI", 20, "bold"))
        self.clock.pack()
        self.date = tk.Label(mid, bg=HEADER, fg=MUTED, font=("Segoe UI", 10))
        self.date.pack()
        tk.Label(row, text="IMI Inventory Tracker", bg=HEADER, fg="#FFFFFF",
                 font=("Segoe UI", 21, "bold")).grid(row=0, column=2, sticky="e")
        bar = tk.Canvas(box, height=3, highlightthickness=0, bg=HEADER)
        bar.pack(fill="x")
        stops = ["#49CCF9", "#7B68EE", "#FD71AF"]

        def draw(_=None):
            bar.delete("all")
            w = bar.winfo_width()
            for x in range(0, w, 4):
                t = x / max(w, 1) * 2
                i = min(int(t), 1)
                bar.create_rectangle(x, 0, x + 4, 3, fill=blend(stops[i + 1], stops[i], t - i), outline="")
        bar.bind("<Configure>", draw)

    def _actions(self):
        row = tk.Frame(self, bg=BG)
        row.pack(fill="x", padx=32, pady=(20, 0))
        self.act = {}
        cmds = [self.show_find, self.show_create, self._soon, self._soon]
        for i, ((label, color), cmd) in enumerate(zip(ACTIONS, cmds)):
            row.columnconfigure(i, weight=1, uniform="a")
            cell = tk.Frame(row, bg=DIS_BD, highlightthickness=1, highlightbackground=DIS_BD)
            cell.grid(row=0, column=i, sticky="ew", padx=6 if 0 < i < 3 else 0)
            bar = tk.Frame(cell, height=4, bg=color)
            bar.pack(fill="x")
            b = tk.Button(cell, text=label, command=cmd, relief="flat", bd=0, font=("Segoe UI", 11, "bold"),
                          height=2, bg=PANEL2, fg=TXT, activebackground="#31355A", activeforeground=TXT,
                          disabledforeground=DIS_TX, cursor="hand2")
            b.pack(fill="both", expand=True)
            b.bind("<Enter>", lambda e, b=b: b["state"] == "normal" and b.config(bg="#31355A"))
            b.bind("<Leave>", lambda e, b=b: b.config(bg=PANEL2 if b["state"] == "normal" else DIS_BG))
            self.act[label] = (b, bar, color, cell, i == 0)
        self._lock_buttons()

    def _lock_buttons(self):
        for label, (b, bar, color, cell, always) in self.act.items():
            on = always or self.user is not None
            b.config(state="normal" if on else "disabled", bg=PANEL2 if on else DIS_BG,
                     cursor="hand2" if on else "X_cursor", text=label if on else f"🔒  {label}")
            bar.config(bg=color if on else blend(color, BG, 0.35))
            cell.config(highlightbackground=BORDER if on else DIS_BD)

    def _activity(self):
        box = tk.Frame(self, bg=BG)
        box.pack(fill="x", padx=32, pady=(0, 20))
        tk.Label(box, text="RECENT ACTIVITY", bg=BG, fg=MUTED, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0, 6))
        wrap = tk.Frame(box, bg=PANEL, highlightthickness=1, highlightbackground=BORDER)
        wrap.pack(fill="x")
        self.log = tk.Text(wrap, height=4, bg=PANEL, fg=TXT, relief="flat", font=("Segoe UI", 10), state="disabled",
                           tabs=("170", "230"), padx=12, pady=4, spacing1=4, spacing3=4, cursor="arrow")
        sb = ttk.Scrollbar(wrap, command=self.log.yview)
        self.log.config(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.log.pack(side="left", fill="x", expand=True)
        self.log.tag_config("time", foreground=MUTED)
        for i, c in enumerate(INI_COLORS):
            self.log.tag_config(f"i{i}", foreground=c, font=("Segoe UI", 10, "bold"))

    def refresh_activity(self):
        rows = self.guarded(self.db.activity) or []
        self.log.config(state="normal")
        self.log.delete("1.0", "end")
        for r in rows:
            t = dt.datetime.strptime(r["timestamp"], FMT).strftime("%b %d, %I:%M %p")
            tag = f"i{sum(map(ord, r['initials'])) % len(INI_COLORS)}"
            self.log.insert("end", t, "time")
            self.log.insert("end", "\t")
            self.log.insert("end", r["initials"], tag)
            self.log.insert("end", f"\t{r['message']}\n")
        self.log.config(state="disabled")

    def _tick(self):
        n = dt.datetime.now()
        self.clock.config(text=n.strftime("%I:%M %p").lstrip("0"))
        self.date.config(text=n.strftime("%A, %B %d, %Y").replace(" 0", " "))
        self.after(1000, self._tick)

    def _soon(self):
        messagebox.showinfo("Coming soon", "This function hasn't been built yet.")

    # -- sign in
    def sign_in(self):
        if getattr(self, "_signing", False):
            return
        self._signing = True
        try:
            self._sign_in()
        finally:
            self._signing = False

    def _sign_in(self):
        ini = self.ini_var.get().strip()
        if not ini:
            self.user = None
        elif not self.user or self.user["initials"] != ini:
            found = self.guarded(self.db.find_user, ini)
            if found is None and self.guarded(lambda: True):
                dlg = RegisterDialog(self, ini)
                self.wait_window(dlg)
                if dlg.result:
                    v = dlg.result
                    final = self.guarded(self.db.add_user, ini, v["first"], v["last"], v["email"])
                    found = self.guarded(self.db.find_user, final) if final else None
                    if found:
                        self.ini_var.set(found["initials"])
                    self.refresh_activity()
            self.user = found
            if not found:
                self.ini_var.set("")
        if self.user:
            name = f"{self.user['first_name']} {self.user['last_name']}"
            self.chip.config(text=f"✓ Signed in as {name}", bg="#183A1F", fg="#8EE07A")
            self.ini.config(highlightbackground="#6BC950")
            self.not_link.config(text=f"Not {name}?")
        else:
            self.chip.config(text="🔒 Required to create or edit data", bg="#3A3115", fg="#FFD54A")
            self.ini.config(highlightbackground=YELLOW)
            self.not_link.config(text="")
            cw = self.content_widget
            if isinstance(cw, ProductCard) and cw.mode != "view":
                self.show_find()
        self._lock_buttons()
        cw = self.content_widget
        if isinstance(cw, FindView):
            cw.redraw()
        elif isinstance(cw, ProductCard) and cw.mode == "view":
            cw.refresh_header()

    def switch_user(self):
        dlg = RegisterDialog(self)
        self.wait_window(dlg)
        if not dlg.result:
            return
        v = dlg.result
        final = self.guarded(self.db.add_user, v["initials"], v["first"], v["last"], v["email"])
        if not final:
            return
        self.refresh_activity()
        self.user = None
        self.ini_var.set(final)
        self.sign_in()
        if final != v["initials"].upper():
            messagebox.showinfo("Initials already in use",
                                f"{v['initials'].upper()} was already taken, so your initials are {final}.")

    # -- views
    def _swap(self, widget):
        if self.content_widget:
            self.content_widget.destroy()
        self.content_widget = widget
        widget.pack(fill="both", expand=True)

    def show_find(self):
        self._swap(FindView(self.content, self))

    def show_item(self, number, mode="view"):
        data = self.guarded(self.db.find_data)
        if not data:
            return
        items, lots = data
        it = next((i for i in items if i["item_number"] == number), None)
        if it:
            mode = mode if (mode == "view" or self.user) else "view"
            self._swap(ProductCard(self.content, self, item=it, mode=mode, lots=lots.get(number, [])))

    def show_create(self):
        if self.user:
            self._swap(ProductCard(self.content, self))


ROW_H = 34
GREEN, GREY = "#2E8B47", "#565B70"


def clip(text, px, font):
    if font.measure(text) <= px:
        return text
    while text and font.measure(text + "…") > px:
        text = text[:-1]
    return text + "…"


class FindView(tk.Frame):
    COLS = [("item_number", "Item Number", 150), ("name", "Name", 0), ("category", "Category", 220),
            ("quantity", "Quantity", 120), ("location", "Location", 160), ("status", "Status", 110), ("edit", "", 64)]

    def __init__(self, master, app):
        super().__init__(master, bg=PANEL, padx=24, pady=20)
        self.app = app
        self.items, self.lots, self.hay, self.rows, self.xs = [], {}, {}, [], []
        self.sort, self.asc, self.total, self.hl = "item_number", True, 0, None
        self.font = tkfont.Font(family="Segoe UI", size=10)
        self.bold = tkfont.Font(family="Segoe UI", size=10, weight="bold")
        tk.Label(self, text="Find Item", bg=PANEL, fg=TXT, font=("Segoe UI", 16, "bold")).pack(anchor="w")
        bar = tk.Frame(self, bg=PANEL)
        bar.pack(fill="x", pady=(12, 4))
        self.q = tk.StringVar()
        e = mk_entry(bar, self.q)
        e.pack(side="left", fill="x", expand=True, ipady=6)
        e.focus_set()
        self.q.trace_add("write", lambda *a: self.fill())
        mk_button(bar, "Refresh", self.load).pack(side="left", padx=(10, 0))
        tk.Label(self, text="Searches item number, name, manufacturer, description/notes, category and lot codes.",
                 bg=PANEL, fg=MUTED, font=("Segoe UI", 9)).pack(anchor="w", pady=(0, 10))
        grid = tk.Frame(self, bg=PANEL)
        grid.pack(fill="both", expand=True)
        grid.columnconfigure(0, weight=1)
        grid.rowconfigure(1, weight=1)
        self.head = tk.Canvas(grid, height=ROW_H, bg=PANEL2, highlightthickness=0)
        self.head.grid(row=0, column=0, sticky="ew")
        self.body = tk.Canvas(grid, bg=PANEL, highlightthickness=0, yscrollincrement=ROW_H)
        self.body.grid(row=1, column=0, sticky="nsew")
        sb = ttk.Scrollbar(grid, command=self.body.yview)
        sb.grid(row=1, column=1, sticky="ns")
        self.body.config(yscrollcommand=sb.set)
        self.empty = tk.Label(self, text="", bg=PANEL, fg=MUTED, font=FONT)
        self.empty.pack(pady=6)
        self.body.bind("<Configure>", self.redraw)
        self.body.bind("<Motion>", self.on_motion)
        self.body.bind("<Leave>", lambda e: self.on_motion(None))
        self.body.bind("<ButtonRelease-1>", self.on_click)
        self.body.bind("<MouseWheel>", lambda e: self.body.yview_scroll(int(-e.delta / 120) or (-1 if e.delta > 0 else 1), "units"))
        self.body.bind("<Button-4>", lambda e: self.body.yview_scroll(-1, "units"))
        self.body.bind("<Button-5>", lambda e: self.body.yview_scroll(1, "units"))
        self.head.bind("<Button-1>", self.on_head)
        self.head.bind("<Motion>", lambda e: self.head.config(cursor="hand2" if self.col_at(e.x) not in (None, "edit") else ""))
        self.load()

    # -- data
    def load(self):
        data = self.app.guarded(self.app.db.find_data)
        self.items, self.lots = data if data else ([], {})
        self.hay = {i["item_number"]: " ".join(str(x or "") for x in (
            i["item_number"], i["name"], i["manufacturer"], i["notes"], i["category"],
            " ".join(self.lots.get(i["item_number"], [])))).casefold() for i in self.items}
        self.fill()

    def skey(self, i):
        k = self.sort
        if k == "quantity":
            q = i["quantity"]
            return (q is None, q if q is not None else 0)
        return str(status_of(i) if k == "status" else (i[k] or "")).casefold()

    def fill(self):
        toks = self.q.get().casefold().split()
        hits = [i for i in self.items if all(t in self.hay[i["item_number"]] for t in toks)]
        self.rows = sorted(hits, key=self.skey, reverse=not self.asc)
        self.empty.config(text="" if hits else ("No items match your search." if self.items else "No items yet."))
        self.body.yview_moveto(0)
        self.redraw()

    # -- drawing
    def layout(self):
        w = max(self.body.winfo_width(), 700)
        name_w = max(180, w - sum(c[2] for c in self.COLS))
        x, self.xs = 0, []
        for key, _, cw in self.COLS:
            cw = name_w if key == "name" else cw
            self.xs.append((x, x + cw))
            x += cw
        self.total = x

    def redraw(self, *_):
        self.layout()
        h, b = self.head, self.body
        h.delete("all")
        b.delete("all")
        for (key, label, _), (x0, x1) in zip(self.COLS, self.xs):
            arrow = (" ▲" if self.asc else " ▼") if key == self.sort else ""
            h.create_text(x0 + 12, ROW_H // 2, text=label + arrow, anchor="w", fill=TXT, font=self.bold)
        self.hl = b.create_rectangle(-5, -5, -4, -4, fill="#23263A", outline="")
        signed = bool(self.app.user)
        for r, i in enumerate(self.rows):
            y, mid = r * ROW_H, r * ROW_H + ROW_H // 2
            archived = status_of(i) == "Archived"
            fg = DIS_TX if archived else TXT
            b.create_line(0, y + ROW_H - 1, self.total, y + ROW_H - 1, fill="#2A2D3B")
            qty = "" if i["quantity"] is None else f"{i['quantity']:g} {i['uom'] or ''}".strip()
            vals = {"item_number": i["item_number"], "name": i["name"], "category": i["category"],
                    "quantity": qty, "location": i["location"] or ""}
            for (key, _, _), (x0, x1) in zip(self.COLS, self.xs):
                if key == "status":
                    b.create_rectangle(x0 + 10, y + 6, x1 - 10, y + ROW_H - 7, fill=GREY if archived else GREEN, outline="")
                    b.create_text((x0 + x1) // 2, mid, text="Archived" if archived else "Active", fill="#FFFFFF", font=self.bold)
                elif key == "edit":
                    if signed:
                        b.create_text((x0 + x1) // 2, mid, text="✎", fill=FOCUS, font=("Segoe UI", 13))
                else:
                    b.create_text(x0 + 12, mid, text=clip(str(vals[key]), x1 - x0 - 20, self.font),
                                  anchor="w", fill=fg, font=self.font)
        b.config(scrollregion=(0, 0, self.total, max(len(self.rows) * ROW_H, 1)))

    # -- interaction
    def col_at(self, x):
        return next((c[0] for c, (x0, x1) in zip(self.COLS, self.xs) if x0 <= x < x1), None)

    def on_head(self, e):
        key = self.col_at(e.x)
        if key in (None, "edit"):
            return
        self.asc = not self.asc if key == self.sort else True
        self.sort = key
        self.fill()

    def row_at(self, e):
        r = int(self.body.canvasy(e.y) // ROW_H)
        return r if 0 <= r < len(self.rows) else None

    def on_motion(self, e):
        r = self.row_at(e) if e else None
        if self.hl is None:
            return
        if r is None:
            self.body.coords(self.hl, -5, -5, -4, -4)
            self.body.config(cursor="")
        else:
            self.body.coords(self.hl, 0, r * ROW_H, self.total, (r + 1) * ROW_H - 1)
            self.body.config(cursor="hand2")

    def on_click(self, e):
        r = self.row_at(e)
        if r is None:
            return
        pencil = self.app.user and e.x >= self.xs[-1][0]
        self.app.show_item(self.rows[r]["item_number"], "edit" if pencil else "view")


class RevDialog(tk.Toplevel):
    def __init__(self, master, number, cur):
        super().__init__(master, bg=PANEL, padx=28, pady=22)
        self.title("Roll revision")
        self.transient(master)
        self.resizable(False, False)
        self.result, self.number = None, number
        i = REV_SEQ.index(cur) if cur else 0
        cands = REV_SEQ[i + 1:i + 7]
        self.rev, self.mode = tk.StringVar(value=cands[0]), tk.StringVar(value="new")
        tk.Label(self, text=f"Roll {number} to a new revision", bg=PANEL, fg=TXT, font=("Segoe UI", 14, "bold")).grid(columnspan=2, sticky="w")
        tk.Label(self, text=f"Current revision: {cur or '(none)'}.  Roll when form, fit or function changes.",
                 bg=PANEL, fg=MUTED, font=FONT).grid(columnspan=2, sticky="w", pady=(0, 14))
        tk.Label(self, text="New revision", bg=PANEL, fg=TXT, font=BOLD).grid(row=2, column=0, sticky="w")
        cb = ttk.Combobox(self, textvariable=self.rev, values=cands, state="readonly", width=8, font=FONT)
        cb.grid(row=2, column=1, sticky="w", padx=(14, 0))
        cb.bind("<<ComboboxSelected>>", lambda e: self.texts())
        self.rb = []
        for r, val in enumerate(("new", "convert"), 3):
            b = tk.Radiobutton(self, variable=self.mode, value=val, bg=PANEL, fg=TXT, selectcolor=FIELD,
                               activebackground=PANEL, activeforeground=TXT, font=FONT, justify="left",
                               wraplength=470, anchor="w")
            b.grid(row=r, column=0, columnspan=2, sticky="w", pady=(12 if r == 3 else 4, 0))
            self.rb.append(b)
        self.texts()
        bar = tk.Frame(self, bg=PANEL)
        bar.grid(columnspan=2, sticky="e", pady=(18, 0))
        mk_button(bar, "Cancel", self.destroy).pack(side="left", padx=8)
        mk_button(bar, "Roll Revision", self.ok, primary=True).pack(side="left")
        self.grab_set()
        self.wait_visibility()
        self.focus_force()

    def new_number(self):
        d = NUM_RE.match(self.number)
        return build_number(d["cat"], d["seq"], d["sub"] or "", self.rev.get(), d["dash"] or "")

    def texts(self):
        new = self.new_number()
        self.rb[0].config(text=f"Create {new} as a separate item. {self.number} stays in stock, so both revisions can be stocked at once.")
        self.rb[1].config(text=f"Convert all stock to {new}. {self.number} is renamed and its stock becomes {new}.")

    def ok(self):
        self.result = {"rev": self.rev.get(), "mode": self.mode.get()}
        self.destroy()


class ProductCard(tk.Frame):
    def __init__(self, master, app, item=None, mode="create", lots=()):
        super().__init__(master, bg=PANEL, padx=28, pady=20)
        self.app, self.item, self.mode = app, item, mode
        lookups = app.guarded(app.db.lookups)
        self.uoms, self.locs = lookups if lookups else (DEFAULT_UOMS, ["Virtual Intake"])
        v = self.v = {k: tk.StringVar() for k in ("name", "cat", "num", "qty", "uom", "loc", "mfr", "gen_n", "min", "sub", "rev", "dash", "link")}
        self.override, self.gen, self.low = tk.BooleanVar(), tk.BooleanVar(), tk.BooleanVar()
        self.track = tk.StringVar(value="NONE")
        v["gen_n"].set("ALL")
        v["uom"].set(self.uoms[0])
        hdr = tk.Frame(self, bg=PANEL)
        hdr.grid(sticky="ew", columnspan=2)
        self.title_lbl = tk.Label(hdr, bg=PANEL, fg=TXT, font=("Segoe UI", 16, "bold"))
        self.title_lbl.pack(side="left")
        self.badge = tk.Label(hdr, bg=PANEL, font=("Segoe UI", 9, "bold"), padx=10, pady=2)
        self.pencil = tk.Button(hdr, text="✎  Edit", command=self.enter_edit, relief="flat", bd=0, font=BOLD,
                                padx=14, pady=4, bg=PANEL2, fg=TXT, activebackground="#31355A",
                                activeforeground=TXT, cursor="hand2")
        sc = tk.Canvas(self, bg=PANEL, highlightthickness=0, yscrollincrement=24)
        sc.grid(sticky="nsew", pady=(10, 0))
        ssb = ttk.Scrollbar(self, command=sc.yview)
        ssb.grid(row=1, column=1, sticky="ns", pady=(10, 0))
        sc.config(yscrollcommand=ssb.set)
        f = self.form = tk.Frame(sc, bg=PANEL)
        win = sc.create_window(0, 0, window=f, anchor="nw")
        f.bind("<Configure>", lambda e: sc.config(scrollregion=sc.bbox("all")))
        sc.bind("<Configure>", lambda e: sc.itemconfig(win, width=e.width - 14))
        self.bind_all("<MouseWheel>", lambda e: sc.yview_scroll(int(-e.delta / 120) or (-1 if e.delta > 0 else 1), "units"))
        self.bind_all("<Button-4>", lambda e: sc.yview_scroll(-1, "units"))
        self.bind_all("<Button-5>", lambda e: sc.yview_scroll(1, "units"))
        self.rowconfigure(1, weight=1)
        self.columnconfigure(0, weight=1)
        f.columnconfigure((0, 1), weight=1, uniform="c")

        def label(parent, text, req=False):
            l = tk.Label(parent, text=text, bg=PANEL, fg=TXT, font=BOLD)
            return l

        def cell(r, c, text, req=False, span=1):
            fr = tk.Frame(f, bg=PANEL)
            fr.grid(row=r, column=c, columnspan=span, sticky="ew", padx=(0, 20) if c == 0 and span == 1 else 0, pady=6)
            lb = label(fr, text + (" *" if req else ""), req)
            lb.pack(anchor="w")
            return fr, lb

        fr, _ = cell(0, 0, "Name", True, 2)
        self.w_name = mk_entry(fr, v["name"])
        self.w_name.pack(fill="x", ipady=6, pady=(4, 0))
        fr, _ = cell(1, 0, "Category", True)
        cat = ttk.Combobox(fr, textvariable=v["cat"], values=CATEGORIES, state="readonly", font=FONT)
        cat.pack(fill="x", pady=(4, 0))
        cat.bind("<<ComboboxSelected>>", lambda e: self.on_category())
        self.cat_note = tk.Label(fr, text="", bg=PANEL, fg="#FFD54A", font=("Segoe UI", 9))
        self.cat_note.pack(anchor="w")
        fr, _ = cell(1, 1, "Item Number", True)
        r = tk.Frame(fr, bg=PANEL)
        r.pack(fill="x", pady=(4, 0))
        self.num = mk_entry(r, v["num"], state="disabled")
        self.num.pack(side="left", fill="x", expand=True, ipady=6)
        tk.Label(r, text="Override", bg=PANEL, fg=MUTED, font=("Segoe UI", 9)).pack(side="left", padx=(12, 6))
        self.ovr_t = Toggle(r, self.override, self.on_override)
        self.ovr_t.pack(side="left")
        self.num_hint = tk.Label(fr, text="", bg=PANEL, fg=MUTED, font=("Segoe UI", 9))
        self.num_hint.pack(anchor="w")
        det = tk.Frame(f, bg=PANEL)
        det.grid(row=2, column=0, columnspan=2, sticky="ew", pady=6)
        self.det, self.det_title = {}, {}
        self.nolink = tk.BooleanVar()

        def dcell(key, title, make):
            dfr = tk.Frame(det, bg=PANEL)
            self.det_title[key] = tk.Label(dfr, text=title, bg=PANEL, fg=TXT, font=BOLD)
            self.det_title[key].pack(anchor="w")
            w = make(dfr)
            w.pack(anchor="w", pady=(4, 0))
            self.det[key] = dfr
            return w
        dcell("sub", "Subcategory *", lambda p: ttk.Combobox(
            p, textvariable=v["sub"], values=[f"{k} – {d}" for k, d in SUBCATS.items()],
            state="readonly", width=38, font=FONT))
        def mk_link(p):
            row = tk.Frame(p, bg=PANEL)
            self.link_cb = ttk.Combobox(row, textvariable=v["link"], state="readonly", width=40, font=FONT)
            self.link_cb.pack(side="left")
            self.nolink_fr = tk.Frame(row, bg=PANEL)
            tk.Label(self.nolink_fr, text="Do not link", bg=PANEL, fg=MUTED, font=("Segoe UI", 9)).pack(side="left", padx=(14, 6))
            Toggle(self.nolink_fr, self.nolink, self.on_nolink).pack(side="left")
            return row
        dcell("link", "Linked PCB", mk_link)
        self.rev_cb = dcell("rev", "Revision", lambda p: ttk.Combobox(
            p, textvariable=v["rev"], width=8, state="readonly", font=FONT))
        dcell("dash", "Dash Number", lambda p: ttk.Combobox(
            p, textvariable=v["dash"], values=["(none)"] + [f"{i:02d}" for i in range(1, 10)], width=8, font=FONT))
        fr, _ = cell(3, 0, "Quantity")
        r = tk.Frame(fr, bg=PANEL)
        r.pack(fill="x", pady=(4, 0))
        mk_entry(r, v["qty"], width=12).pack(side="left", ipady=6)
        ttk.Combobox(r, textvariable=v["uom"], values=self.uoms, font=FONT).pack(side="left", fill="x", expand=True, padx=(8, 0))
        tk.Label(fr, text="Type a new unit in the box to add one.", bg=PANEL, fg=MUTED, font=("Segoe UI", 9)).pack(anchor="w")
        fr, self.loc_lbl = cell(3, 1, "Location")
        ttk.Combobox(fr, textvariable=v["loc"], values=self.locs, state="readonly", font=FONT).pack(fill="x", pady=(4, 0))
        fr, _ = cell(4, 0, "Manufacturer")
        self.w_mfr = mk_entry(fr, v["mfr"])
        self.w_mfr.pack(fill="x", ipady=6, pady=(4, 0))
        fr, _ = cell(4, 1, "Tracking")
        r = tk.Frame(fr, bg=PANEL)
        r.pack(fill="x", pady=(4, 0))
        self.radios = []
        for txt, val in (("None", "NONE"), ("LOT", "LOT"), ("SN", "SN")):
            rb = tk.Radiobutton(r, text=txt, variable=self.track, value=val, bg=PANEL, fg=TXT, selectcolor=FIELD,
                                activebackground=PANEL, activeforeground=TXT, disabledforeground=DIS_TX, font=FONT,
                                command=self.sync_tracking)
            rb.pack(side="left", padx=(0, 12))
            self.radios.append(rb)
        tk.Label(r, text="Generate", bg=PANEL, fg=MUTED, font=("Segoe UI", 9)).pack(side="left", padx=(8, 6))
        self.gen_t = Toggle(r, self.gen, self.sync_tracking)
        self.gen_t.pack(side="left")
        self.gen_n = ttk.Combobox(r, textvariable=v["gen_n"], values=["ALL"], width=6, font=FONT)
        fr, _ = cell(5, 0, "Low stock notification")
        r = tk.Frame(fr, bg=PANEL)
        r.pack(fill="x", pady=(4, 0))
        self.low_t = Toggle(r, self.low, self.sync_low)
        self.low_t.pack(side="left")
        self.min_lbl = tk.Label(r, text="Alert when below", bg=PANEL, fg=MUTED, font=FONT)
        self.min_e = mk_entry(r, v["min"], width=10)
        fr, _ = cell(6, 0, "Description / Notes", span=2)
        self.notes = tk.Text(fr, height=5, bg=FIELD, fg=TXT, insertbackground=TXT, relief="flat", font=FONT,
                             highlightthickness=2, highlightbackground=BORDER, highlightcolor=FOCUS, wrap="word")
        self.notes.pack(fill="x", pady=(4, 0))
        meta = tk.Frame(f, bg=PANEL2)
        meta.grid(row=7, column=0, columnspan=2, sticky="ew", pady=(10, 0))
        if item:
            texts = [f"Created by {item['created_by']} · {fmt_ts(item['created_at'])}",
                     f"Last edited by {item['edited_by']} · {fmt_ts(item['edited_at'])}"]
        else:
            texts = ["Created by / on:  set when saved", "Last edited by / on:  set when saved"]
        for t in texts:
            tk.Label(meta, text=t, bg=PANEL2, fg=MUTED, font=("Segoe UI", 10), padx=12, pady=6).pack(side="left")
        if lots:
            tk.Label(f, text="Lot codes: " + ", ".join(lots), bg=PANEL, fg=TXT, font=FONT, wraplength=900,
                     justify="left").grid(row=8, column=0, columnspan=2, sticky="w", pady=(8, 0))
        self.err = tk.Label(self, text="", bg=PANEL, fg=ERR, font=BOLD, anchor="w")
        self.err.grid(sticky="ew", pady=(8, 0), columnspan=2)
        self.bar = tk.Frame(self, bg=PANEL)
        self.bar.grid(sticky="e", pady=(6, 0), columnspan=2)
        creating = mode == "create"
        self.items = (app.guarded(app.db.items) or []) if creating else []
        self.existing = {_norm(i["item_number"]) for i in self.items}
        self.seq, self._forced, self._dup_shown = None, False, False
        if creating:
            v["num"].trace_add("write", lambda *a: self.check_dup())
            for k in ("sub", "rev"):
                v[k].trace_add("write", lambda *a: self.preview())
            v["dash"].trace_add("write", lambda *a: self.on_dash())
            v["link"].trace_add("write", lambda *a: self.refresh_seq())
            self.on_category()
        else:
            self.populate(item)
        v["qty"].trace_add("write", lambda *a: self.sync_qty())
        self.sync_qty()
        self.sync_tracking()
        self.sync_low()
        if not creating:
            self.num_hint.config(text="")
            self.lock_all()
        if mode == "edit":
            self.enter_edit()
        else:
            self.refresh_header()
            self.build_buttons()

    def destroy(self):
        for ev in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            self.unbind_all(ev)
        super().destroy()

    # -- view / edit modes
    def status(self):
        return status_of(self.item) if self.item else "Active"

    def _walk(self, w):
        for c in w.winfo_children():
            yield c
            yield from self._walk(c)

    def lock_all(self):
        for w in self._walk(self.form):
            if isinstance(w, ttk.Combobox):
                w.config(state="disabled")
            elif isinstance(w, (tk.Entry, tk.Radiobutton)):
                w.config(state="disabled")
            elif isinstance(w, tk.Text):
                w.config(state="disabled", bg=DIS_BG)
            elif isinstance(w, Toggle):
                w.set_enabled(False)

    def unlock(self):
        """Edit mode: only name, description, manufacturer, tracking option and stock notification."""
        for w in (self.w_name, self.w_mfr, self.min_e):
            w.config(state="normal")
        self.notes.config(state="normal", bg=FIELD)
        for rb in self.radios:
            rb.config(state="normal")
        self.low_t.set_enabled(True)

    def populate(self, it):
        v = self.v
        v["name"].set(it["name"] or "")
        v["cat"].set(it["category"] or "")
        v["num"].set(it["item_number"])
        v["uom"].set(it["uom"] or "")
        if it["quantity"] is not None:
            v["qty"].set(f"{it['quantity']:g}")
        v["loc"].set(it["location"] or "")
        v["mfr"].set(it["manufacturer"] or "")
        self.notes.insert("1.0", it["notes"] or "")
        self.track.set(it["tracking"] or "NONE")
        self.low.set(it["low_stock_alert"] == "Yes")
        if it["min_quantity"] is not None:
            v["min"].set(f"{it['min_quantity']:g}")
        shown = []
        if it.get("subcategory"):
            v["sub"].set(f"{it['subcategory']} – {SUBCATS.get(it['subcategory'], '')}")
            shown.append("sub")
        if it.get("revision"):
            v["rev"].set(it["revision"])
            shown.append("rev")
        if it.get("dash_number"):
            v["dash"].set(it["dash_number"])
            shown.append("dash")
        for k in shown:
            self.det[k].pack(side="left", padx=(0, 20))

    def refresh_header(self):
        m = self.mode
        self.title_lbl.config(text="New Product Card" if m == "create" else
                              ("Edit " if m == "edit" else "") + self.item["item_number"])
        self.badge.pack_forget()
        self.pencil.pack_forget()
        if m != "create":
            archived = self.status() == "Archived"
            self.badge.config(text="ARCHIVED" if archived else "ACTIVE", bg=GREY if archived else GREEN, fg="#FFFFFF")
            self.badge.pack(side="left", padx=14)
            if m == "view" and self.app.user:
                self.pencil.pack(side="right")

    def build_buttons(self):
        for w in self.bar.winfo_children():
            w.destroy()
        m, app = self.mode, self.app

        def add(text, cmd, primary=False):
            mk_button(self.bar, text, cmd, primary).pack(side="left", padx=8)
        if m == "create":
            add("Cancel", app.show_find)
            add("Save Item", self.save, True)
        elif m == "view":
            add("Back to Find", app.show_find)
        else:
            num, archived = self.item["item_number"], self.status() == "Archived"
            add("Cancel", lambda: app.show_item(num, "view"))
            add("Reactivate Item" if archived else "Archive Item", self.toggle_archive)
            if not archived and self.rollable():
                add("Roll Revision…", self.roll_rev)
            if not archived:
                add("Save Changes", self.save_edit, True)

    def enter_edit(self):
        if not self.app.user:
            return
        self.mode = "edit"
        self.lock_all()
        if self.status() == "Active":
            self.unlock()
        self.sync_low()
        self.refresh_header()
        self.build_buttons()

    def rollable(self):
        m = NUM_RE.match(str(self.item["item_number"]))
        return bool(m) and m["cat"] in REV_TRACKED | REV_OPTIONAL

    def save_edit(self):
        v = {k: x.get().strip() for k, x in self.v.items()}
        if not v["name"]:
            self.err.config(text="Name is required.")
            return
        min_q = None
        if self.low.get():
            try:
                min_q = float(v["min"])
            except ValueError:
                self.err.config(text="Enter the minimum quantity that should trigger the low stock alert.")
                return
        num = self.item["item_number"]
        changes = dict(name=v["name"], manufacturer=v["mfr"] or None, notes=self.notes.get("1.0", "end").strip() or None,
                       tracking=self.track.get(), low_stock_alert="Yes" if self.low.get() else "No", min_quantity=min_q)
        if self.app.guarded(self.app.db.update_item, num, changes, self.app.user["initials"]):
            self.app.refresh_activity()
            self.app.show_item(num, "view")

    def toggle_archive(self):
        num, archived = self.item["item_number"], self.status() == "Archived"
        q = (f"Reactivate {num}?" if archived else
             f"Archive {num}?\n\nIt stays in the system (part numbers are never deleted) and can be reactivated later.")
        if not messagebox.askyesno("Reactivate item" if archived else "Archive item", q):
            return
        if self.app.guarded(self.app.db.set_status, num, "Active" if archived else "Archived", self.app.user["initials"]):
            self.app.refresh_activity()
            self.app.show_item(num, "view")

    def roll_rev(self):
        num = self.item["item_number"]
        dlg = RevDialog(self.app, num, NUM_RE.match(num)["rev"] or "")
        self.app.wait_window(dlg)
        if not dlg.result:
            return
        try:
            new = self.app.db.roll_revision(num, dlg.result["rev"], dlg.result["mode"], self.app.user["initials"])
        except FileBusy:
            messagebox.showerror("Error", BUSY_MSG)
            return
        except ValueError as e:
            messagebox.showerror("Cannot roll revision", str(e))
            return
        except Exception as e:
            messagebox.showerror("Error", f"Could not access the data file.\n\n{e}")
            return
        self.app.refresh_activity()
        self.app.show_item(new, "view")

    # -- field behaviour
    def check_dup(self):
        n = self.v["num"].get().strip()
        msg = ""
        if n and _norm(n) in self.existing:
            msg = f"Item number {n} already exists. Every item must have a unique number."
        if msg:
            self.err.config(text=msg)
            self._dup_shown = True
        elif self._dup_shown:
            self.err.config(text="")
            self._dup_shown = False

    def code(self):
        return self.v["cat"].get()[:2]

    def on_category(self):
        code = self.code()
        v = self.v
        self.cat_note.config(text="Requires inventory manager approval." if code == "09" else "")
        want = [k for k, on in (("sub", code == "06"), ("link", code in ("01", "02")),
                                ("rev", code in REV_TRACKED | REV_OPTIONAL), ("dash", code in DASH_CATS)) if on]
        for k, fr in self.det.items():
            fr.pack_forget()
        for k in ("sub", "link", "rev", "dash"):
            if k in want:
                self.det[k].pack(side="left", padx=(0, 20))
        if code != "06":
            v["sub"].set("")
        if code not in DASH_CATS:
            v["dash"].set("(none)")
        elif not v["dash"].get():
            v["dash"].set("(none)")
        self.nolink.set(False)
        self.link_cb.config(state="readonly")
        self.nolink_fr.pack_forget()
        links, seen = [], set()
        if code in ("01", "02"):
            seqs01 = set(_seqs(self.items, "01"))
            for i in self.items:
                m = NUM_RE.match(str(i["item_number"]))
                if not m or m["seq"] in seen:
                    continue
                if (code == "02" and m["cat"] == "01") or (code == "01" and m["cat"] == "02" and int(m["seq"]) not in seqs01):
                    seen.add(m["seq"])
                    links.append(f"{i['item_number']}  {i['name']}")
        if code == "02":
            self.det_title["link"].config(text="Linked PCB *")
            self.nolink_fr.pack(side="left")
        elif code == "01":
            self.det_title["link"].config(text="Link to existing PCBA (optional)")
        self.link_cb.config(values=["(none)"] + links)
        v["link"].set("(none)" if code in ("01", "02") else "")
        if code == "02":
            self.rev_cb.config(values=list(REV_LETTERS), state="readonly")
            v["rev"].set("A")
        elif code in REV_TRACKED:
            self.rev_cb.config(values=["A"], state="disabled")
            v["rev"].set("A")
        elif code in REV_OPTIONAL:
            self.rev_cb.config(values=["(none)", "A"], state="readonly")
            v["rev"].set("(none)")
        else:
            v["rev"].set("(none)")
        if code == "00":  # externally numbered (MPN / vendor number): typed by hand
            self._forced = True
            self.override.set(True)
            self.ovr_t.set_enabled(False)
            v["num"].set("")
            self.on_override()
        else:
            self.ovr_t.set_enabled(True)
            if self._forced:
                self._forced = False
                self.override.set(False)
                self.on_override()
        self.refresh_seq()

    def link_seq(self):
        tok = self.v["link"].get().split()
        m = NUM_RE.match(tok[0]) if tok else None
        return int(m["seq"]) if m else None

    def refresh_seq(self):
        code = self.code()
        self.seq = None
        if code and code != "00":
            linked = self.link_seq() if code in ("01", "02") else None
            if linked:
                self.seq = linked
            elif code != "02" or self.nolink.get():  # unlinked: next number nobody shares
                self.seq = self.app.guarded(self.app.db.peek_sequence, code)
        if code == "02" and self.seq:
            self.suggest_rev()
        self.preview()
        self.update_hint()

    def suggest_rev(self):
        v = self.v
        dash = "" if v["dash"].get() in ("", "(none)") else v["dash"].get().zfill(2)
        used = {m["rev"] for i in self.items if (m := NUM_RE.match(str(i["item_number"]))) and m["cat"] == "02"
                and int(m["seq"]) == self.seq and (m["dash"] or "") == dash}
        v["rev"].set(next((c for c in REV_LETTERS if c not in used), "A"))

    def on_dash(self):
        if self.code() == "02" and self.seq:
            self.suggest_rev()
        self.preview()

    def on_nolink(self):
        if self.nolink.get():
            self.v["link"].set("(none)")
        self.link_cb.config(state="disabled" if self.nolink.get() else "readonly")
        self.refresh_seq()

    def update_hint(self):
        code, on = self.code(), self.override.get()
        if code == "00":
            t = "External number: enter the MPN / vendor number exactly."
        elif on:
            t = "Override: must follow NN-AAAA-NNNNNA-NN."
        elif code == "02" and self.seq is None:
            t = "Choose a linked PCB, or turn on Do not link."
        else:
            t = "Assigned from the Part Numbering System."
        self.num_hint.config(text=t)

    def preview(self):
        code, v = self.code(), self.v
        if self.override.get() or code in ("", "00"):
            return
        if self.seq is None:
            v["num"].set("")
            return
        rev = "" if v["rev"].get() in ("", "(none)") else v["rev"].get()
        dash = "" if v["dash"].get() in ("", "(none)") else v["dash"].get().zfill(2)[:2]
        sub = (v["sub"].get()[:4] or "XXXX") if code == "06" else ""
        v["num"].set(build_number(code, self.seq, sub, rev, dash))

    def on_override(self):
        on = self.override.get()
        self.num.config(state="normal" if on else "disabled")
        self.update_hint()
        if not on:
            self.preview()

    def qty(self):
        t = self.v["qty"].get().strip()
        if not t:
            return None
        try:
            q = float(t)
        except ValueError:
            raise ValueError("Quantity must be a number.")
        if q < 0:
            raise ValueError("Quantity cannot be negative.")
        return q

    def sync_qty(self):
        has = bool(self.v["qty"].get().strip())
        self.loc_lbl.config(text="Location" + (" *" if has else ""))
        if not has:
            self.gen.set(False)
        self.sync_tracking()

    def sync_tracking(self):
        can = self.mode == "create" and self.track.get() == "LOT" and bool(self.v["qty"].get().strip())
        if not can:
            self.gen.set(False)
        self.gen_t.set_enabled(can)
        if self.gen.get():
            self.gen_n.pack(side="left", padx=(10, 0))
        else:
            self.gen_n.pack_forget()

    def sync_low(self):
        if self.low.get():
            self.min_lbl.pack(side="left", padx=(14, 6))
            self.min_e.pack(side="left", ipady=4)
        else:
            self.min_lbl.pack_forget()
            self.min_e.pack_forget()

    # -- save
    def save(self):
        v = {k: x.get().strip() for k, x in self.v.items()}
        try:
            if not v["name"]:
                raise ValueError("Name is required.")
            if not v["cat"]:
                raise ValueError("Category is required.")
            if not v["num"]:
                raise ValueError("Item Number is required.")
            code = v["cat"][:2]
            rev = "" if v["rev"] in ("", "(none)") else v["rev"]
            dash = "" if v["dash"] in ("", "(none)") else v["dash"].zfill(2)
            if code == "02" and not self.link_seq() and not self.nolink.get():
                raise ValueError("Choose a linked PCB, or turn on Do not link.")
            if not self.override.get() and code != "00":
                if code == "06" and not v["sub"]:
                    raise ValueError("Choose a subcategory for chemical formulations.")
                if dash and not re.fullmatch(r"\d{2}", dash):
                    raise ValueError("Dash number must be 2 digits, e.g. 01.")
            qty = self.qty()
            if qty is not None and not v["loc"]:
                raise ValueError("Location is required when a quantity is entered.")
            if qty is not None and not v["uom"]:
                raise ValueError("Choose a unit of measure.")
            uom = v["uom"]
            match = next((u for u in self.uoms if _norm(u) == _norm(uom)), None)
            uom = match or re.sub(r"\s+", " ", uom)  # reuse the existing spelling if it's a duplicate
            lots = make_lots(qty, uom, v["gen_n"]) if self.gen.get() else []
            min_q = None
            if self.low.get():
                try:
                    min_q = float(v["min"])
                except ValueError:
                    raise ValueError("Enter the minimum quantity that should trigger the low stock alert.")
        except ValueError as e:
            self.err.config(text=str(e))
            return
        rec = dict(item_number=v["num"], name=v["name"], category=v["cat"], quantity=qty,
                   uom=uom if qty is not None else None, location=v["loc"] or None, manufacturer=v["mfr"] or None,
                   notes=self.notes.get("1.0", "end").strip() or None, tracking=self.track.get(),
                   low_stock_alert="Yes" if self.low.get() else "No", min_quantity=min_q,
                   subcategory=v["sub"][:4], revision=rev, dash_number=dash,
                   link_seq=self.link_seq() if code in ("01", "02") else None)
        try:
            n = self.app.db.save_item(rec, lots, self.app.user["initials"], not self.override.get())
        except FileBusy:
            messagebox.showerror("Error", BUSY_MSG)
            return
        except ValueError as e:
            self.err.config(text=str(e))
            messagebox.showerror("Cannot save item", str(e))
            return
        except Exception as e:
            messagebox.showerror("Error", f"Could not save.\n\n{e}")
            return
        self.app.refresh_activity()
        self.app.show_find()
        messagebox.showinfo("Saved", f"Item {n} created." + (f"\nLot codes: {', '.join(c for c, _ in lots)}" if lots else ""))


if __name__ == "__main__":
    App().mainloop()